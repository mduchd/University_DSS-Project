"""Tạo sổ tay Word về nhiệm vụ ML và toàn bộ luồng DSS của nhóm."""

from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "So_tay_ML_va_tong_quan_de_tai_DSS.docx"
EVAL = json.loads((ROOT / "models" / "evaluation.json").read_text(encoding="utf-8"))


def heading(doc: Document, number: str, title: str, level: int = 1) -> None:
    doc.add_heading(f"{number} {title}", level=level)


def para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def table(doc: Document, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Light Shading Accent 1"
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        grid.rows[0].cells[i].text = value
    for row in rows:
        cells = grid.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    doc.add_paragraph()


def qa(doc: Document, number: str, question: str, answer: str) -> None:
    heading(doc, number, question, 2)
    para(doc, answer)


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.3)
    sec.right_margin = Cm(2.1)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(20, 65, 95)
    title = doc.add_heading("SỔ TAY MACHINE LEARNING VÀ QUY TRÌNH ĐỀ TÀI DSS", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("La Bàn Đại Học — tài liệu để học, bảo vệ đồ án và đối chiếu với mã nguồn")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph("Ảnh chụp hiện trạng mã nguồn: 10/10/2026. File .py và evaluation.json là nguồn đúng khi có thay đổi sau ngày này.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    heading(doc, "1.", "Bài toán, vai trò ML và ranh giới hệ thống")
    heading(doc, "1.1", "Bài toán nghiệp vụ", 2)
    para(doc, "Học sinh THPT nhập tổ hợp và điểm ba môn, chọn sở thích/nhóm ngành/khu vực. Hệ thống trả về các phương án trường–ngành để tham khảo, chia thành An toàn, Phù hợp và Thử sức. Đây là hệ hỗ trợ quyết định (Decision Support System, DSS): hệ thống tổng hợp dữ kiện và giải thích gợi ý, còn học sinh/phụ huynh là người ra quyết định cuối cùng. Không có xác suất đỗ được hiệu chỉnh và không có cam kết trúng tuyển.")
    heading(doc, "1.2", "Ranh giới phần việc của thành viên ML", 2)
    para(doc, "Phần ML dự báo điểm chuẩn trên thang 30 cho năm mục tiêu từ dữ liệu lịch sử, kiểm soát chất lượng dữ liệu, xây feature (đặc trưng), so sánh mô hình, đóng gói predictor và đưa dự báo cho backend. Bộ xếp hạng phương án ở services/recommendation_service.py là phần khuyến nghị của nhóm, không phải đầu ra trực tiếp của XGBoost. Model không nhận điểm của một học sinh để tính xác suất đỗ.")
    para(doc, "Luồng đang chạy dùng điểm chuẩn 2018–2024 để tạo mốc dự báo 2025; nếu thiếu feature để dự báo, ứng dụng dùng điểm chuẩn 2024 làm fallback (giá trị thay thế). Dữ liệu tuyển sinh 2025–2026 do crawler lưu riêng chưa được đưa vào model đang triển khai. Vì vậy demo hiện tại không phải tư vấn theo điểm chuẩn thời gian thực năm 2026.")

    heading(doc, "2.", "Sơ đồ cấu trúc và vai trò các file")
    heading(doc, "2.1", "Cây thư mục thuộc nhiệm vụ ML", 2)
    tree = (
        "models/\n"
        "├─ prepare_cutoff_dataset.py      → sàng lọc nguồn, target, khóa modeling, báo cáo loại\n"
        "├─ analyze_cutoff_step2.py       → rà soát bản ghi loại, map định danh, review queue\n"
        "├─ review_high_priority_identities.py → ưu tiên các cặp trường/ngành cần xác minh\n"
        "├─ build_cutoff_features.py      → lag, thống kê lịch sử, exam t−1, cold-start\n"
        "├─ train_baselines.py            → các mốc tham chiếu và linear regression\n"
        "├─ train_cutoff_model.py         → preprocessing, Ridge/RF/XGBoost, chọn và test\n"
        "├─ predictor.py                  → API dự báo độc lập từ artifact\n"
        "├─ generate_cutoff_evaluation_charts.py → hình đánh giá\n"
        "├─ model.joblib / evaluation.json → pipeline đã fit / báo cáo kết quả\n"
        "└─ CUTOFF_MODEL_LIMITATIONS.md   → phạm vi, cold-start, leakage, hợp đồng output\n"
        "services/cutoff_forecast_service.py → ghép forecast vào phương án và fallback\n"
        "scripts/build_forecast_cache.py   → tạo cache forecast theo batch\n"
        "notebooks/ml/01_…10_*.ipynb     → bản chú giải để nhóm đọc/thử nghiệm; .py là nguồn production\n"
        "tests/test_cutoff_dataset.py, test_cutoff_model.py, test_predictor.py → kiểm thử tự động"
    )
    run = doc.add_paragraph()
    run.add_run(tree).font.name = "Consolas"
    run.style = "No Spacing"
    heading(doc, "2.2", "File dữ liệu/đầu ra quan trọng", 2)
    table(doc, ["Đường dẫn", "Vai trò"], [
        ("data/processed/admission_ml_base.csv", "Dòng đã qua điều kiện nền cho bài toán hồi quy điểm chuẩn."),
        ("data/processed/admission_ml_rejected.csv", "Dòng bị loại kèm lý do, bảo đảm truy vết."),
        ("data/processed/admission_ml_step2.csv", "Dữ liệu sau bước rà soát định danh và mapping."),
        ("data/processed/admission_ml_review_queue.csv", "Danh sách các trường hợp phải xác minh thủ công."),
        ("data/processed/admission_ml_features.csv", "Bảng feature theo năm; nguồn train/validation/test."),
        ("data/processed/admission_ml_cold_start_2024.csv", "Phân tích nhóm test không có lịch sử riêng."),
        ("data/processed/admission_ml_baseline_metrics.csv", "Chỉ số các baseline; xem cả coverage (tỷ lệ có dự báo)."),
        ("data/processed/admission_ml_model_predictions_2024.csv", "Dự báo test 2024 để tính residual và vẽ hình."),
        ("data/processed/admission_ml_feature_importance.csv", "Mức quan trọng feature từ model đã chọn."),
        ("models/model.joblib", "Artifact gồm preprocessing và estimator, chỉ nạp từ nguồn tin cậy."),
        ("models/evaluation.json", "Split, cấu hình ứng viên, metric, tiêu chí chọn model, giới hạn."),
        ("docs/figures/cutoff_model/", "Biểu đồ validation, test, residual, feature importance."),
    ])
    heading(doc, "2.3", "Các hàm chính để đọc code", 2)
    table(doc, ["File", "Hàm/khối quan trọng và điều cần hiểu"], [
        ("prepare_cutoff_dataset.py", "main: kiểm tra schema, năm, thang điểm, phương thức, tổ hợp và khóa trùng; joined_reasons: giữ nhiều lý do loại; atomic_write_*: ghi file an toàn."),
        ("analyze_cutoff_step2.py", "classify_note: phân loại ghi chú; build_institution_map/build_program_map: chuẩn hóa định danh; annotate_rejections và build_review_queue: giải thích dòng cần xem lại."),
        ("review_high_priority_identities.py", "pairwise_institution_review/pairwise_program_review: phát hiện cặp nghi vấn; annotate_high_queue: ưu tiên kiểm tra, không auto-merge."),
        ("build_cutoff_features.py", "build_series_history_features: lag và lịch sử; add_previous_year_aggregate/add_prior_history_aggregate: fallback theo nhóm; add_exam_distribution_features: phổ điểm t−1; add_cold_start_design: phân lớp lạnh; validate_exact_lags: phát hiện leakage."),
        ("train_baselines.py", "fit_linear_regression/predict_linear: nghiệm bình phương tối thiểu; add_group_mean_baseline: mốc cho cold-start; evaluate_predictions: MAE/RMSE/R² và coverage."),
        ("train_cutoff_model.py", "prepare_model_frame: khóa feature hợp lệ; make_preprocessor: impute + OneHot + scale khi cần; fit_validation_candidates: fit ứng viên chỉ trên train; refit_selected: fit lại train+validation; grouped_test_metrics/feature_importance/main: test và xuất artifact."),
        ("predictor.py", "load_artifact: cache pipeline; _derive_features: chuẩn hóa record; CutoffPredictor.predict/predict_batch: kiểm tra đầu vào và dự báo độc lập."),
        ("generate_cutoff_evaluation_charts.py", "plot_candidate_comparison, plot_test_segments, plot_actual_vs_predicted, plot_residual_distribution, plot_feature_importance: tạo hình cho báo cáo."),
        ("cutoff_forecast_service.py", "CutoffForecastService: chuẩn bị feature store và ghép dự báo với bảng phương án; fallback nếu thiếu mapping/feature."),
        ("build_forecast_cache.py", "main: tạo cache dự báo trước, tránh phải dự báo lặp lại cho từng request."),
    ])

    heading(doc, "3.", "Quy trình từ lý thuyết đến thực hành của phần ML")
    heading(doc, "3.1", "Định nghĩa đơn vị dự báo và target", 2)
    para(doc, "Một quan sát là điểm chuẩn năm t của một khóa modeling: năm × mã trường × mã ngành/chương trình tuyển sinh × tổ hợp môn. Target y(s,t) = cutoff_score_30 là điểm chuẩn đã quy về thang 0–30. s chỉ một chuỗi cùng chương trình/tổ hợp qua các năm. Cần thống nhất grain (đơn vị của một dòng) trước khi khử trùng; nếu cùng khóa có nhiều target xung đột thì không lấy trung bình tùy tiện.")
    heading(doc, "3.2", "Làm sạch và truy vết dòng bị loại", 2)
    para(doc, "Đầu vào là bảng admissions đã chuẩn hóa. Code chỉ giữ phương thức xét tuyển THPT phù hợp, năm 2018–2024, target là số trong [0,30], tổ hợp hợp lệ, mã định danh đủ. Dòng trùng hoàn toàn có thể gộp; cùng khóa nhưng điểm khác nhau phải loại/đưa vào review. File rejected và quality_report ghi rõ nguyên nhân để sau này kiểm toán. Không tự biến điểm thiếu thành 0 vì 0 là giá trị có nghĩa khác với chưa biết.")
    heading(doc, "3.3", "Rà soát định danh và cold-start", 2)
    para(doc, "Tên hiển thị thay đổi hoặc mã tái sử dụng có thể làm đứt lịch sử chuỗi. Bước 2 lập map định danh trường/ngành và review queue; bước ưu tiên cao đối chiếu các cặp giống nhau bằng token similarity (độ giống từ) rồi yêu cầu xác minh thủ công. Cold-start nghĩa là một chuỗi s chưa có quan sát trước t. Không thể lấy lag riêng; hệ thống dùng thống kê quá khứ theo trường+tổ hợp, nhóm ngành+tổ hợp, trường, nhóm ngành, tổ hợp rồi toàn cục. Fallback là thông tin ít cá nhân hóa hơn nên phải đo sai số riêng.")
    heading(doc, "3.4", "Feature engineering không rò rỉ thời gian", 2)
    para(doc, "Feature cho năm t chỉ được dùng dữ liệu có sẵn trước t. lag_k(s,t) là điểm của đúng năm t−k nếu tồn tại; historical_mean(s,t) là trung bình các y(s,τ) với τ<t; n_history(s,t) là số quan sát lịch sử; trend là độ dốc điểm theo năm quá khứ. Các thống kê trường/nhóm ngành/tổ hợp phải tính theo dữ liệu quá khứ tương tự. Phổ điểm thi dùng năm t−1 vì sản phẩm đang giả định dự báo trước khi có điểm thi năm t. Nếu đổi sang dự báo sau kỳ thi, phải version lại thời điểm và đánh giá lại.")
    heading(doc, "3.5", "Chia dữ liệu, chọn model và kiểm tra", 2)
    para(doc, "Có ba cách chia thường gặp: chia ngẫu nhiên, cross-validation ngẫu nhiên và chia theo thời gian. Hai cách đầu có thể cho mô hình nhìn thấy thông tin tương lai của cùng trường/ngành khi dự báo năm trước đó. Đồ án chọn train <=2022 (55.713 dòng), validation 2023 (10.524), test 2024 (19.827). Dùng validation để chọn feature/cấu hình; giữ test 2024 cho báo cáo cuối. Sau chọn, fit lại trên train+validation rồi mới test. Không tuning theo test.")
    heading(doc, "3.6", "Baseline và mô hình ứng viên", 2)
    para(doc, "Historical Mean (trung bình lịch sử) và Last Value (điểm gần nhất) là mốc dễ giải thích nhưng thiếu coverage khi chuỗi mới. Group Mean dự báo được cold-start bằng trung bình nhóm; baseline hybrid kết hợp Last Value và Group Mean đạt coverage 100%. Linear Regression/Ridge là mô hình tuyến tính có kiểm soát overfitting; Random Forest là trung bình nhiều cây quyết định; XGBoost là boosting tuần tự sửa sai của các cây trước. Chọn cây vì quan hệ giữa trường, ngành, thời gian và lịch sử điểm chuẩn phi tuyến; vẫn phải thắng baseline trên validation chứ không chọn vì tên thuật toán.")
    heading(doc, "3.7", "Preprocessing và huấn luyện", 2)
    para(doc, "Feature phân loại được điền giá trị thiếu và OneHotEncoder(handle_unknown='ignore') để trường/ngành mới không làm hỏng dự báo. Feature số được SimpleImputer điền median học từ train, kèm missing indicator; Ridge dùng StandardScaler để hệ số cùng thang. Tất cả preprocessing và model cùng Pipeline. Random Forest thử n_estimators, max_depth, min_samples_leaf; XGBoost thử độ sâu, learning_rate, subsample, colsample và early stopping trên validation 2023. Random seed 42 giúp tái lập trong phạm vi môi trường tương đồng.")
    heading(doc, "3.8", "Đánh giá, đóng gói và tích hợp", 2)
    para(doc, "Đánh giá MAE, RMSE, R² trên toàn bộ, known-history, cold-start và nhóm trường/ngành lớn; đo coverage của baseline. Artifact joblib chứa cả preprocessing đã fit và estimator. predictor.py kiểm tra đầu vào, hỗ trợ category mới, trả predicted_cutoff và is_cold_start; service ghép theo khóa để backend dùng. Test tự động kiểm tra khóa không trùng, không leakage, split, predictor load lại, miền dự báo và output không phải 'xác suất đỗ'.")

    heading(doc, "4.", "Công thức và ký hiệu cần giải thích khi vấn đáp")
    table(doc, ["Ký hiệu", "Ý nghĩa trong đề tài"], [
        ("s", "Một chuỗi trường–ngành/chương trình–tổ hợp."),
        ("t, τ", "Năm cần dự báo; τ là một năm lịch sử với τ<t."),
        ("y(s,t), ŷ(s,t)", "Điểm chuẩn thực tế và điểm chuẩn dự báo của s trong năm t."),
        ("N", "Số quan sát có target trong tập đang đánh giá; không phải số năm."),
        ("u, i", "Hồ sơ học sinh u và phương án trường–ngành i trong recommender."),
        ("m, n", "Nếu viết ma trận user–item: m = số người dùng, n = số phương án. Dự án chưa có ma trận tương tác đa người dùng để train collaborative filtering."),
    ])
    para(doc, "Lag: Lₖ(s,t) = y(s,t−k) nếu điểm năm t−k tồn tại; thiếu thì để missing. Mean lịch sử: μ(s,t) = [Σ_{τ<t} y(s,τ)] / |{τ<t có quan sát}|. Mọi τ trong tổng đều nhỏ hơn t; đây là điều kiện chống leakage cốt lõi.")
    para(doc, "MAE = (1/N)Σⱼ|yⱼ−ŷⱼ|: sai số tuyệt đối trung bình, cùng đơn vị 'điểm'. RMSE = √[(1/N)Σⱼ(yⱼ−ŷⱼ)²]: phạt mạnh sai số lớn. R² = 1 − [Σⱼ(yⱼ−ŷⱼ)² / Σⱼ(yⱼ−ȳ)²]: mức giải thích biến thiên; có thể âm nếu tệ hơn dự báo hằng số trung bình. j đánh số dòng trong tập đánh giá; ȳ là điểm thực tế trung bình của đúng tập ấy.")
    para(doc, "Ridge tìm hệ số β bằng cách tối thiểu hóa Σⱼ(yⱼ−xⱼᵀβ)² + α||β||₂². xⱼ là vector feature của dòng j; α điều khiển mức phạt hệ số lớn. Random Forest lấy trung bình dự báo của nhiều cây, mỗi cây học trên mẫu/feature ngẫu nhiên. XGBoost xây dần các cây để giảm loss hồi quy có regularization; learning_rate thu nhỏ đóng góp cây mới, early stopping dừng khi MAE validation không cải thiện.")
    para(doc, "Trong recommender hiện hành, utility f(u,i) là điểm phù hợp theo nội dung: 100 × [w₁·admission_fit(u,i) + w₂·interest_fit(u,i) + w₃·salary_norm(i) + w₄·demand_norm(i)], với trọng số do quy tắc hồ sơ chọn và tổng trọng số bằng 1. Đây là weighted scoring (chấm điểm có trọng số), không phải cosine similarity được học từ tương tác. Các ngưỡng lọc và tính ổn định lịch sử tiếp tục ảnh hưởng kết quả trong mã; đọc services/recommendation_service.py để thấy quy tắc cụ thể.")

    heading(doc, "5.", "Bảng kết quả thực nghiệm và nhận xét")
    heading(doc, "5.1", "Baseline: chú ý coverage", 2)
    baseline_rows = []
    for name, values in EVAL["models"]["baselines"].items():
        va, te = values["validation"], values["test"]
        baseline_rows.append((name, f"{va['mae']:.3f}", f"{te['mae']:.3f}", f"{te['coverage']*100:.1f}%"))
    table(doc, ["Baseline", "MAE val 2023", "MAE test 2024", "Coverage test"], baseline_rows)
    heading(doc, "5.2", "Ứng viên ML và model được chọn", 2)
    candidate_rows = []
    for item in EVAL["models"]["validation_candidates"]:
        metric = item["validation"]["all"]
        candidate_rows.append((item["name"], f"{metric['mae']:.3f}", f"{metric['rmse']:.3f}", f"{metric['r2']:.3f}", f"{item['fit_seconds']:.1f}"))
    table(doc, ["Ứng viên", "MAE val", "RMSE val", "R² val", "Fit (giây)"], candidate_rows)
    test = EVAL["final_test_2024"]
    para(doc, f"Model chọn: {EVAL['selection']['selected_candidate']} theo MAE validation. Test 2024: MAE {test['all']['mae']:.3f}, RMSE {test['all']['rmse']:.3f}, R² {test['all']['r2']:.3f}; known-history MAE {test['known_history']['mae']:.3f}, cold-start MAE {test['cold_start']['mae']:.3f}. Nhóm cold-start chiếm 11.572/19.827 dòng = 58,36% nên sai số toàn bộ bị ảnh hưởng đáng kể. Không được so MAE baseline chỉ có coverage 41,6% như thể nó dự báo đủ 100%.")
    para(doc, "CPU/GPU và thời gian: code thiết lập XGBoost tree_method='hist', Random Forest n_jobs=-1 và không cấu hình GPU. Có thể nói đường chạy được cấu hình theo CPU; không có log phần cứng để khẳng định máy cụ thể. evaluation.json lưu thời gian fit từng ứng viên trên validation (ví dụ XGBoost_1 khoảng 65,6 giây), không ghi tổng thời gian toàn pipeline hay dung lượng RAM đỉnh. Muốn báo cáo tài nguyên chính xác phải benchmark lại trên máy được nêu rõ CPU/RAM/phiên bản thư viện.")

    heading(doc, "6.", "Quy trình toàn đề tài và vai trò các thành viên khác")
    heading(doc, "6.1", "Từ ba nguồn đến một quyết định", 2)
    para(doc, "(1) Thu thập Admission (điểm chuẩn và mô tả tuyển sinh), Exam (phổ điểm tổng hợp), VietJobs (tin tuyển dụng). (2) Mỗi luồng làm sạch/chuẩn hóa riêng bằng scripts/standardize_admission.py, standardize_exam.py và standardize_vietjobs.py; loại dữ liệu cá nhân khỏi phần công khai. (3) scripts/build_unified_dataset.py và mapping ngành–nghề tạo master_admission.csv cùng các bảng nghề/lương/nhu cầu. (4) Nhánh ML tạo dự báo điểm chuẩn; nhánh recommender nhận hồ sơ học sinh, lọc ứng viên, ghép dự báo và nội dung nghề nghiệp, tính utility, chia ba bucket. (5) Flask app.py cung cấp API; frontend HTML/CSS/JS hiển thị gợi ý và lưu nguyện vọng ở localStorage. (6) Tests và các file báo cáo kiểm tra toàn luồng.")
    heading(doc, "6.2", "Những lựa chọn kỹ thuật của phần khuyến nghị", 2)
    para(doc, "Lọc cộng tác (Collaborative Filtering) cần lịch sử nhiều người dùng tương tác với nhiều trường/ngành; hệ thống hiện chưa có dữ liệu đó nên không thể học sở thích tập thể một cách đáng tin. Lọc dựa nội dung (Content-Based Filtering) dùng đặc trưng của hồ sơ và phương án: tổ hợp, ngưỡng điểm, sở thích-ngành, vùng, lương tham khảo, nhu cầu tuyển dụng. Hiện code tính interest fit bằng so khớp từ khóa và weighted scoring, KHÔNG dùng cosine similarity. AHP/TOPSIS được giữ ở module tham khảo nhưng không nằm trên đường API chính. Phổ điểm Exam là feature bối cảnh t−1 của ML, không phải điểm cá nhân của người dùng.")
    heading(doc, "6.3", "Sơ đồ ra quyết định bằng lời", 2)
    para(doc, "Hồ sơ u → kiểm tra điểm/tổ hợp → lọc master theo điều kiện cứng → gắn predicted_cutoff 2025 hoặc fallback lịch sử 2024 → tính gap = tổng điểm u − mốc điểm → loại phương án ngoài ngưỡng thử sức → tính admission_fit, interest_fit, lương và nhu cầu đã chuẩn hóa → tổng hợp weighted score → xếp hạng từng bucket → trả danh sách có giải thích → học sinh quyết định cuối. API hiện giới hạn tối đa 15 mục mỗi bucket và có top 5 lựa chọn nổi bật; không có bước cố định 'cosine lấy top 100 rồi hiện 10'.")

    heading(doc, "7.", "Mười câu hỏi vấn đáp trọng tâm phần ML")
    questions = [
        ("Vì sao đây là regression chứ không phải classification?", "Target là điểm chuẩn liên tục trên thang 30; dự báo một con số. Ba bucket An toàn/Phù hợp/Thử sức được quyết định ở tầng khuyến nghị sau khi so điểm học sinh với mốc, không phải nhãn train của XGBoost."),
        ("Một dòng dữ liệu đại diện cho cái gì?", "Một mốc điểm chuẩn của đúng năm, trường, mã ngành/chương trình và tổ hợp THPT. Nếu grain không rõ, gộp các mốc khác nhau sẽ tạo nhãn sai."),
        ("Thế nào là data leakage trong đề tài?", "Feature năm t vô tình chứa điểm chuẩn hay phổ điểm chỉ biết ở t hoặc sau t. Ví dụ tính mean trên toàn bộ 2018–2024 trước khi chia tập sẽ làm 2024 nhìn thấy chính nó."),
        ("Vì sao chia theo năm mà không random split?", "Sản phẩm dự báo tương lai từ quá khứ. Chia ngẫu nhiên có thể đưa 2024 vào train trong khi đánh giá 2023, làm metric lạc quan giả."),
        ("Cold-start được xử lý ra sao?", "Khi chuỗi trường–ngành–tổ hợp chưa có quá khứ, lag là missing. Sử dụng thống kê nhóm quá khứ và missing indicator; báo cáo metric cold-start riêng thay vì gán 0."),
        ("Vì sao phải có baseline?", "Để chứng minh mô hình phức tạp thực sự tốt hơn cách dự báo đơn giản có cùng phạm vi mẫu. Baseline không có coverage đầy đủ phải được ghi rõ."),
        ("Vì sao chọn XGBoost?", "Trong candidate đã thử, XGBoost_1 có MAE validation toàn bộ thấp nhất (1,360). Nó mô tả quan hệ phi tuyến và xử lý tổ hợp feature; quyết định không dựa vào test 2024."),
        ("MAE, RMSE, R² nói gì và có thể gây hiểu sai ở đâu?", "MAE dễ diễn giải theo điểm; RMSE nhạy lỗi lớn; R² so với dự báo trung bình. Metric toàn bộ che lấp chênh lệch known/cold-start và không phản ánh coverage nếu baseline bỏ mẫu."),
        ("Vì sao lưu cả preprocessing trong joblib?", "Vì category encoding, imputation và scale phải dùng đúng tham số học khi train; tách estimator khỏi preprocessing sẽ lệch feature giữa train và inference."),
        ("Có thể gọi predicted_cutoff là xác suất đỗ không?", "Không. Đây là dự báo điểm chuẩn, chưa hiệu chỉnh theo phân phối điểm học sinh, chỉ tiêu, nguyện vọng và cơ chế xét tuyển. Backend chỉ được trả mốc điểm và cảnh báo tham khảo."),
    ]
    for idx, (q, a) in enumerate(questions, 1):
        qa(doc, f"7.{idx}", q, a)

    heading(doc, "8.", "Giải đáp ảnh câu hỏi vấn đáp về DSS/recommender")
    questions2 = [
        ("Quy trình ra quyết định và người ra quyết định là ai?", "Hệ thống biến dữ liệu nguồn + hồ sơ u thành danh sách phương án có điểm, ngưỡng và bối cảnh nghề. Người ra quyết định cuối là học sinh, có thể cùng phụ huynh/cố vấn; mục tiêu là tìm phương án phù hợp năng lực/sở thích/rủi ro, không tối ưu một mục tiêu duy nhất."),
        ("Đây là loại quyết định gì?", "Quyết định bán cấu trúc, nhiều tiêu chí và có bất định: một phần lọc/tính điểm theo quy tắc và model, phần còn lại là sở thích, tài chính, hoàn cảnh cá nhân và lựa chọn của người học."),
        ("Thuật toán nào có giám sát, thuật toán nào không?", "Ridge, Random Forest và XGBoost dự báo điểm chuẩn là supervised learning vì có nhãn y là điểm chuẩn lịch sử. Bộ xếp hạng content-based hiện tại là quy tắc có trọng số, không phải một mô hình ranking được train giám sát. Nó cũng không phải collaborative filtering."),
        ("Ba thành phần của một hệ khuyến nghị là gì?", "User u là học sinh/hồ sơ; item i là phương án trường–ngành–tổ hợp; interaction thông thường là xem, lưu, đăng ký, chọn. Repo chỉ lưu nguyện vọng cục bộ trên trình duyệt, chưa có interaction log đa người dùng để học ma trận m×n."),
        ("f(u,i)→R nghĩa là gì; u và i là gì?", "f gán điểm utility R cho cặp hồ sơ học sinh u và phương án i. Trong app hiện tại R là recommendation_score được tính từ độ khớp điểm, sở thích, lương và nhu cầu. Điểm càng cao thì phương án càng được ưu tiên trong bucket."),
        ("Có những bộ lọc khuyến nghị nào và nhóm dùng loại nào?", "Thông dụng: content-based, collaborative, knowledge/rule-based và hybrid. Nhóm dùng content-based kết hợp lọc điều kiện cứng + weighted scoring + mốc ML; chưa dùng collaborative do thiếu tương tác. Đây là một kiến trúc lai theo thành phần, nhưng không nên gọi là hybrid CF nếu không có CF."),
        ("Có dùng cosine không? Nếu nói u,i thì vector là gì?", "Không trong luồng recommender hiện tại. Nếu tương lai triển khai cosine thì phải định nghĩa vector cùng không gian đặc trưng, chuẩn hóa và trọng số; vector u biểu diễn sở thích/năng lực học sinh, vector i biểu diễn thuộc tính ngành/trường. Hiện interest_fit dựa trên khớp từ khóa, không phải tích vô hướng cosine."),
        ("Có lấy top 100 từ cosine rồi hiện 10 không?", "Không. app.py có limit tra cứu mặc định 15, tối đa 100 cho API search; recommendation_service.py xếp hạng và trả tối đa 15 phương án trong mỗi bucket, thêm top 5 nổi bật. Không nên gộp nhầm quy tắc search và recommendation."),
        ("Hàm mục tiêu, loss và m,n cần nói sao?", "ML: N là số dòng được đánh giá; xⱼ là vector feature, yⱼ là điểm thực, ŷⱼ là dự báo; loss bình phương/MAE để fit/chọn model. Recommender: f(u,i) là utility, không được huấn luyện từ nhãn tương tác; nếu viết ma trận tương tác giả định thì m là số user và n là số item, nhưng dự án chưa dùng ma trận ấy."),
        ("Từ ảnh yêu cầu mô tả bài toán bằng đoạn văn thế nào?", "Bài toán của nhóm là hỗ trợ học sinh u chọn trong n phương án i. Với mỗi i, hệ thống có điểm chuẩn lịch sử, trường/ngành, tổ hợp và bối cảnh việc làm; hồ sơ u có điểm ba môn, tổ hợp và sở thích. Nhánh ML dùng N quan sát điểm chuẩn lịch sử để dự báo mốc của i ở năm mục tiêu; nhánh khuyến nghị lọc điều kiện và chấm f(u,i) để sắp xếp. n ở đây là số phương án ứng viên của một phiên; m chỉ xuất hiện nếu ta giả định tập nhiều người dùng, không phải biến bắt buộc của mô hình hiện hành."),
    ]
    for idx, (q, a) in enumerate(questions2, 1):
        qa(doc, f"8.{idx}", q, a)

    heading(doc, "9.", "Cách đọc lại code và tái lập kết quả")
    para(doc, "Mở notebooks/ml theo thứ tự 01→10 để đọc từng hàm với vai trò ở ngay trước cell mã. Notebook không tự gọi main(); đây là bản học tập, không thay thế pipeline production. Nếu sửa .py, chạy python scripts/generate_ml_notebooks.py để đồng bộ. Lệnh kiểm tra: python scripts/validate_ml_notebooks.py và python -m pytest -q. Muốn tái tạo từ dữ liệu, phải chạy các bước chuẩn hóa nguồn trước rồi tuần tự prepare_cutoff_dataset.py → analyze_cutoff_step2.py → review_high_priority_identities.py → build_cutoff_features.py → train_baselines.py → train_cutoff_model.py → generate_cutoff_evaluation_charts.py → build_forecast_cache.py. Các bước này ghi file output; nên chạy trên bản sao dữ liệu hoặc xác nhận trạng thái git trước khi thực hành.")
    para(doc, "Nguồn đối chiếu chính: README.md; PRODUCT.md; docs/data-model-report.md; models/CUTOFF_MODEL_LIMITATIONS.md; models/evaluation.json; services/recommendation_service.py; app.py và các tests. Số liệu trong Word là từ artifact đã lưu, không phải benchmark mới ở lượt tạo tài liệu này.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
